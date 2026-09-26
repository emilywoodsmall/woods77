"""
odometer_to_excel.py
Reads odometer readings from an OdoLogger over Bluetooth and saves them to an Excel file.

Setup:
    pip install -r requirements.txt

Usage:
    1. Pair your computer with "OdoLogger_BUS_01" in Bluetooth settings.
    2. Find the serial port:
         Mac:     /dev/cu.OdoLogger_BUS_01      (run:  ls /dev/cu.*)
         Windows: the "Outgoing" COM port in Device Manager > Ports (e.g. COM5)
    3. Run:
         python odometer_to_excel.py /dev/cu.OdoLogger_BUS_01
       or try it without any hardware:
         python odometer_to_excel.py --test

Lines from the logger look like:
    ODO,BUS_01,123456.125,J1939-HR,27.60,00     (vehicle, km, source, battery V, source address)
    ERR,BUS_01,NO_DATA  /  ERR,BUS_01,STALE
"""

import argparse
import datetime
import os
import random
import time

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import serial

KM_TO_MILES = 0.621371

LOG_HEADERS = ["Date", "Time", "Vehicle", "Odometer (km)", "Odometer (miles)",
               "Miles since last", "Source", "Battery (V)"]
LOG_WIDTHS = [12, 10, 14, 15, 17, 17, 18, 12]
TOTALS_HEADERS = ["Vehicle", "First reading (mi)", "Last reading (mi)", "Total miles",
                  "First date", "Last date"]

# Column numbers in the Log sheet (1-based, like Excel)
COL_DATE, COL_VEHICLE, COL_KM, COL_MILES = 1, 3, 4, 5


def make_new_excel(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(LOG_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate(LOG_WIDTHS):
        ws.column_dimensions[get_column_letter(i + 1)].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Totals")
    ws2.append(TOTALS_HEADERS)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for i in range(len(TOTALS_HEADERS)):
        ws2.column_dimensions[get_column_letter(i + 1)].width = 18

    wb.save(path)
    print("made new excel file: " + path)


def find_last_km(ws, vehicle):
    """Return the most recent odometer (km) logged for this vehicle, or None."""
    for row in range(ws.max_row, 1, -1):
        if ws.cell(row=row, column=COL_VEHICLE).value == vehicle:
            return ws.cell(row=row, column=COL_KM).value
    return None


def update_totals(wb):
    ws = wb["Log"]
    ws2 = wb["Totals"]
    if ws2.max_row > 1:
        ws2.delete_rows(2, ws2.max_row)

    first, last = {}, {}
    for row in range(2, ws.max_row + 1):
        vehicle = ws.cell(row=row, column=COL_VEHICLE).value
        if vehicle is None:
            continue
        miles = ws.cell(row=row, column=COL_MILES).value
        date = ws.cell(row=row, column=COL_DATE).value
        if vehicle not in first:
            first[vehicle] = (miles, date)
        last[vehicle] = (miles, date)

    for vehicle in first:
        first_mi, first_date = first[vehicle]
        last_mi, last_date = last[vehicle]
        ws2.append([vehicle, first_mi, last_mi, round(last_mi - first_mi, 1), first_date, last_date])


def save_reading(path, vehicle, km, source, volts):
    if not os.path.exists(path):
        make_new_excel(path)

    try:
        wb = openpyxl.load_workbook(path)
    except PermissionError:
        print("CAN'T OPEN EXCEL FILE - close it in Excel and try again")
        return

    ws = wb["Log"]
    last_km = find_last_km(ws, vehicle)

    if last_km is not None:
        if km == last_km:
            print(vehicle + " didn't move, not saving")
            return
        if km < last_km:
            # an odometer should never go backwards: wrong module, bad frame, or swapped logger
            print("WARNING: %s odometer went backwards (%.3f -> %.3f km), not saving"
                  % (vehicle, last_km, km))
            return

    miles = round(km * KM_TO_MILES, 1)
    since_last = 0 if last_km is None else round((km - last_km) * KM_TO_MILES, 1)

    now = datetime.datetime.now()
    date = now.strftime("%Y-%m-%d")
    t = now.strftime("%H:%M:%S")

    ws.append([date, t, vehicle, km, miles, since_last, source, volts])
    update_totals(wb)

    try:
        wb.save(path)
        print("saved: %s %s  %s  %.1f mi  (+%.1f)" % (date, t, vehicle, miles, since_last))
    except PermissionError:
        print("CAN'T SAVE - close the Excel file first. This reading was lost.")


def handle_line(line, path):
    parts = [p.strip() for p in line.split(",")]

    if parts[0] == "ODO" and len(parts) >= 3:
        vehicle = parts[1]
        try:
            km = float(parts[2])
        except ValueError:
            print("bad number: " + line)
            return

        source = parts[3] if len(parts) >= 4 else ""
        if len(parts) >= 6 and parts[5]:
            source += " (SA " + parts[5] + ")"

        volts = None
        if len(parts) >= 5 and parts[4]:
            try:
                volts = float(parts[4])
            except ValueError:
                pass

        save_reading(path, vehicle, km, source, volts)

    elif parts[0] == "ERR":
        reason = parts[2] if len(parts) >= 3 else "?"
        vehicle = parts[1] if len(parts) >= 2 else "?"
        if reason == "NO_DATA":
            print(vehicle + ": no odometer seen yet (key off? wrong bitrate? wiring?)")
        elif reason == "STALE":
            print(vehicle + ": odometer stopped updating (key turned off?)")
        else:
            print(vehicle + ": logger error: " + line)

    else:
        print("got something unexpected: " + line)


def run_test_mode(path):
    print("TEST MODE - making fake readings (Ctrl+C to stop)")
    fake_km = 80467.200
    while True:
        fake_km += random.randint(1, 30)
        volts = round(random.uniform(26.8, 28.2), 2)
        handle_line("ODO,TEST_BUS,%.3f,J1939-HR,%.2f,00" % (fake_km, volts), path)
        time.sleep(3)


def run_bluetooth(port, path):
    while True:
        try:
            print("connecting to " + port + " ...")
            with serial.Serial(port, 115200, timeout=5) as bt:
                print("connected!")
                bt.write(b"READ\n")   # ask for a reading right away
                while True:
                    line = bt.readline().decode("utf-8", errors="ignore").strip()
                    if line:
                        handle_line(line, path)
        except serial.SerialException as e:
            print("lost connection or couldn't connect: " + str(e))
            print("trying again in 10 seconds...")
            time.sleep(10)


def main():
    parser = argparse.ArgumentParser(description="Log OdoLogger odometer readings to Excel.")
    parser.add_argument("port", nargs="?", help="Bluetooth serial port, e.g. /dev/cu.OdoLogger_BUS_01 or COM5")
    parser.add_argument("--file", default="fleet_odometer_log.xlsx", help="Excel file to write")
    parser.add_argument("--test", action="store_true", help="make up readings (no hardware needed)")
    args = parser.parse_args()

    print("OdoLogger excel saver")
    try:
        if args.test:
            run_test_mode(args.file)
        elif args.port:
            run_bluetooth(args.port, args.file)
        else:
            parser.error("give a serial port, or use --test")
    except KeyboardInterrupt:
        print("bye")


if __name__ == "__main__":
    main()
